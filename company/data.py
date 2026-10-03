"""Bounded datasets, reusable steps and typed runtime substitution."""
import copy
import csv
import json
import re
from pathlib import Path

TOKEN = re.compile(r'\$\{([A-Za-z_][A-Za-z_0-9]*)\}')

def substitute(value, variables):
    if isinstance(value, str):
        match = TOKEN.fullmatch(value)
        if match:
            return copy.deepcopy(variables[match[1]])
        return TOKEN.sub(lambda m: str(variables[m[1]]), value)
    if isinstance(value, list):
        return [substitute(x, variables) for x in value]
    if isinstance(value, dict):
        return {k: substitute(v, variables) for k, v in value.items()}
    return value

def rows(source, directory):
    if isinstance(source, str):
        root = Path(directory).resolve()
        path = (root / source).resolve()
        if not path.is_relative_to(root):
            raise ValueError('Dataset must be inside configuration directory')
        if path.stat().st_size > 2_000_000:
            raise ValueError('Dataset exceeds 2 MB')
        if path.suffix.lower() == '.xlsx':
            from openpyxl import load_workbook
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                iterator = book.active.iter_rows(values_only=True)
                headers = next(iterator)
                if any(not isinstance(h, str) for h in headers) or len(set(headers)) != len(headers):
                    raise ValueError('Excel needs unique text column names')
                source = []
                for row in iterator:
                    source.append(dict(zip(headers, row)))
                    if len(source) > 1000:
                        raise ValueError('Dataset exceeds 1000 rows')
            finally:
                book.close()
        else:
            with path.open(encoding='utf-8-sig', newline='') as f:
                source = list(csv.DictReader(f)) if path.suffix.lower() == '.csv' else json.load(f)
    if not isinstance(source, list) or not 1 <= len(source) <= 1000 or any(not isinstance(r, dict) for r in source):
        raise ValueError('Dataset requires 1–1000 object rows')
    return source

def expand(config, directory):
    blocks = config.get('blocks', {})
    def steps(items, stack=()):
        result = []
        for item in items:
            if 'use' in item:
                name = item['use']
                if name in stack or len(stack) >= 16:
                    raise ValueError('Recursive step block')
                result.extend(steps(blocks[name], stack + (name,)))
            else:
                result.append(copy.deepcopy(item))
            if len(result) > 500:
                raise ValueError('Too many steps')
        return result
    result = []
    for test in config['tests']:
        dataset = rows(config['datasets'][test['dataset']], directory) if 'dataset' in test else [{}]
        for index, row in enumerate(dataset):
            item = copy.deepcopy(test)
            item['name'] += f' [row {index + 1}]' if 'dataset' in test else ''
            item['variables'] = {**test.get('variables', {}), **row}
            for key in ('setup', 'steps', 'cleanup'):
                item[key] = steps(test.get(key, []))
            if not any(s.get('action') in ('assert', 'expect_text', 'expect_count', 'expect_url') or 'expect_status' in s or 'expect_json' in s for s in item['steps']):
                raise ValueError('Every test requires an explicit expected outcome')
            result.append(item)
            if len(result) > 1000:
                raise ValueError('Too many expanded tests')
    if not result:
        raise ValueError('No tests configured')
    return result
